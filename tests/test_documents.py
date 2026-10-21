import io
import zipfile

import pytest
from docx import Document

from privacy_gateway.documents import DocumentProcessor
from privacy_gateway.errors import PrivacyError


def pdf_bytes(blank=False, image=False):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    if not blank:
        page.insert_text((72, 72), "Patient: Mira Vale\nPhone: 9876543210")
    if image:
        pix = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 8, 8))
        pix.clear_with(255)
        page.insert_image(pymupdf.Rect(20, 20, 40, 40), pixmap=pix)
    result = doc.tobytes()
    doc.close()
    return result


def docx_bytes():
    doc = Document()
    doc.add_paragraph("Patient: Mira Vale")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Phone: 9876543210"
    doc.sections[0].header.paragraphs[0].text = "Email: header@example.test"
    doc.sections[0].footer.paragraphs[0].text = "MRN-45821"
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


@pytest.mark.parametrize(
    "filename,content",
    [
        ("synthetic.txt", b"Patient: Mira Vale\nPhone: 9876543210"),
        ("synthetic.json", b'{"patient_name":"Mira Vale","phone":9876543210}'),
    ],
)
def test_text_documents(client, filename, content):
    response = client.post("/mask/document", files={"file": (filename, content)})
    assert response.status_code == 200, response.text
    assert "Mira Vale" not in response.text and "9876543210" not in response.text


def test_searchable_pdf(client):
    response = client.post("/mask/document", files={"file": ("test.pdf", pdf_bytes())})
    assert response.status_code == 200, response.text
    assert "Mira Vale" not in response.text and "9876543210" not in response.text


def test_docx_table_header_footer(client):
    response = client.post("/mask/document", files={"file": ("test.docx", docx_bytes())})
    assert response.status_code == 200, response.text
    assert all(
        value not in response.text
        for value in ["Mira Vale", "9876543210", "header@example.test", "MRN-45821"]
    )


@pytest.mark.parametrize(
    "filename,content",
    [
        ("bad.txt", b"\xff\xfe"),
        ("bad.pdf", b"invalid"),
        ("bad.docx", b"invalid"),
        ("bad.exe", b"data"),
        ("empty.txt", b""),
        ("nul.txt", b"abc\x00secret"),
        ("bad.json", b'{"a":"first","a":"lost"}'),
        ("bad.json", b'{"a":NaN}'),
    ],
)
def test_unreadable_documents_rejected(client, filename, content):
    response = client.post("/mask/document", files={"file": (filename, content)})
    assert response.status_code in {422, 415}


def test_blank_pdf_page_rejected():
    with pytest.raises(PrivacyError):
        DocumentProcessor().extract(pdf_bytes(blank=True), "scan.pdf")


def test_mixed_image_pdf_rejected():
    with pytest.raises(PrivacyError, match="partial_extraction_risk"):
        DocumentProcessor().extract(pdf_bytes(image=True), "mixed.pdf")


def test_partial_multipage_pdf_rejected():
    import pymupdf

    with pymupdf.open(stream=pdf_bytes(), filetype="pdf") as doc:
        doc.new_page()
        content = doc.tobytes()
    with pytest.raises(PrivacyError):
        DocumentProcessor().extract(content, "partial.pdf")


def test_docx_unsupported_part_rejected():
    output = io.BytesIO(docx_bytes())
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr("word/embeddings/secret.bin", "synthetic")
    with pytest.raises(PrivacyError, match="partial_extraction_risk"):
        DocumentProcessor().extract(output.getvalue(), "embedded.docx")


def pdf_with_attachments(*attachments):
    """attachments: (name, c2pa_markers) pairs."""
    import pymupdf

    with pymupdf.open(stream=pdf_bytes(), filetype="pdf") as doc:
        for name, c2pa in attachments:
            if not c2pa:
                doc.embfile_add(name, b"synthetic attachment", filename=name)
                continue
            # Indirect file spec and typed stream, as C2PA-signing tools write them.
            stream = doc.get_new_xref()
            doc.update_object(stream, "<</Type/EmbeddedFile/Subtype/application#2Fc2pa>>")
            doc.update_stream(stream, b"synthetic manifest")
            spec = doc.get_new_xref()
            doc.update_object(
                spec,
                f"<</Type/Filespec/AFRelationship/C2PA_Manifest/F({name})/UF({name})"
                f"/EF<</F {stream} 0 R>>>>",
            )
            doc.xref_set_key(
                doc.pdf_catalog(), "Names/EmbeddedFiles/Names", f"[({name}) {spec} 0 R]"
            )
        return doc.tobytes()


def test_pdf_content_credentials_manifest_is_ignored():
    content = pdf_with_attachments(("Content Credentials", True))
    assert "Mira Vale" in DocumentProcessor().extract(content, "signed.pdf")


@pytest.mark.parametrize(
    "attachments",
    [
        [("Content Credentials", False)],
        [("Content Credentials", True), ("notes.txt", False)],
    ],
)
def test_pdf_other_attachments_rejected(attachments):
    with pytest.raises(PrivacyError, match="unsupported_pdf"):
        DocumentProcessor().extract(pdf_with_attachments(*attachments), "attached.pdf")


# --- Layout: reading order and structure ----------------------------------------------------

LEFT_COLUMN = (
    "Masking clinical text before it reaches a hosted model is a common requirement. "
    "Our cohort was followed for twelve weeks and every record was reviewed twice."
)
RIGHT_COLUMN = (
    "The masked pipeline kept most of its task accuracy while removing direct identifiers. "
    "Residual risk remains for rare diagnoses and is discussed below in detail."
)


def two_column_pdf(pages=1):
    import pymupdf

    doc = pymupdf.open()
    for number in range(1, pages + 1):
        page = doc.new_page()
        if number == 1:
            page.insert_text((72, 60), "Layout-Aware Masking", fontsize=18)
        page.insert_textbox(pymupdf.Rect(72, 100, 300, 300), LEFT_COLUMN, fontsize=9)
        page.insert_textbox(pymupdf.Rect(312, 100, 540, 300), RIGHT_COLUMN, fontsize=9)
        page.insert_text((72, 815), f"Journal of Synthetic Studies - page {number}", fontsize=8)
    result = doc.tobytes()
    doc.close()
    return result


def words(text):
    return " ".join(text.split())


def test_two_column_pdf_is_read_one_column_at_a_time():
    text = DocumentProcessor().extract(two_column_pdf(), "paper.pdf")
    flat = words(text)
    # Each column comes out whole and in order, not spliced line by line.
    assert words(LEFT_COLUMN) in flat and words(RIGHT_COLUMN) in flat
    assert flat.index(words(LEFT_COLUMN)) < flat.index(words(RIGHT_COLUMN))
    assert text.startswith("# Layout-Aware Masking\n")


def test_pages_are_marked_and_running_footers_labelled():
    text = DocumentProcessor().extract(two_column_pdf(pages=2), "paper.pdf")
    assert text.startswith("--- page 1 ---\n") and "\n--- page 2 ---\n" in text
    assert "Page footer: Journal of Synthetic Studies - page 1" in text
    assert "Page footer: Journal of Synthetic Studies - page 2" in text


def test_label_and_value_on_one_row_stay_together():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # Short side-by-side text is a label column, not two text columns.
    for row, (label, value) in enumerate(
        [("Client contact", "Mira Vale"), ("Phone", "9876543210"), ("Ward", "7B")]
    ):
        page.insert_text((72, 100 + 20 * row), label, fontsize=10)
        page.insert_text((330, 100 + 20 * row), value, fontsize=10)
    content = doc.tobytes()
    doc.close()
    lines = DocumentProcessor().extract(content, "form.pdf").splitlines()
    assert [words(line) for line in lines] == [
        "Client contact Mira Vale",
        "Phone 9876543210",
        "Ward 7B",
    ]
    # The gap between them is written wide, so a labeled value stops before the next column.
    assert "   " in lines[0]


def test_bordered_pdf_table_becomes_rows_of_cells():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    rows = [("Method", "Accuracy", "Reviewer"), ("Ours", "94%", "Mira Vale")]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(72 + c * 150, 100 + r * 18, 222 + c * 150, 118 + r * 18)
            page.draw_rect(rect, color=(0, 0, 0), width=0.5)
            page.insert_text((rect.x0 + 4, rect.y1 - 5), value, fontsize=9)
    content = doc.tobytes()
    doc.close()
    text = DocumentProcessor().extract(content, "table.pdf")
    assert (
        "| Method | Accuracy | Reviewer |\n| --- | --- | --- |\n| Ours | 94% | Mira Vale |" in text
    )


def structured_docx():
    doc = Document()
    doc.add_heading("Discharge Summary", level=1)
    doc.add_paragraph("Admitted with suspected pneumonia.")
    for item in ("Metformin 500 mg twice daily", "Insulin 10 units at night"):
        doc.add_paragraph(item, style="List Number")
    table = doc.add_table(rows=2, cols=3)
    for r, row in enumerate([("Name", "Phone", "Ward"), ("Mira Vale", "9876543210", "7B")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    paragraph = doc.add_paragraph("Follow-up with")
    paragraph.add_run().add_break()
    paragraph.add_run("the ward team")
    paragraph.add_run("\tweekly")
    doc.sections[0].header.paragraphs[0].text = "Synthetic record"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_docx_keeps_headings_lists_tables_and_breaks():
    text = DocumentProcessor().extract(structured_docx(), "summary.docx")
    assert text.splitlines()[0] == "Page header: Synthetic record"
    assert "# Discharge Summary\n" in text
    assert "1. Metformin 500 mg twice daily\n2. Insulin 10 units at night" in text
    assert "| Name | Phone | Ward |\n| --- | --- | --- |\n| Mira Vale | 9876543210 | 7B |" in text
    # Breaks and tabs no longer glue words together ("withthe", "teamweekly").
    assert "Follow-up with\nthe ward team\tweekly" in text


def test_table_label_columns_are_masked_by_rules(client):
    response = client.post("/mask/document", files={"file": ("summary.docx", structured_docx())})
    assert response.status_code == 200, response.text
    masked = response.json()["sanitized_text"]
    assert "Mira Vale" not in masked and "9876543210" not in masked
    # Table structure survives masking.
    assert "| [PERSON_001] | [PHONE_001] | 7B |" in masked
